"""Build the project presentation: Pulse_Monitor_Presentation.pptx

Uses python-pptx (and the images produced by make_diagrams.py + screenshots
from the running system). Re-run any time the assets change.

    pip install python-pptx
    python build_presentation.py
"""
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from lxml import etree

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

HERE = Path(__file__).resolve().parent
ASSETS = HERE / "assets"
OUT = HERE / "Pulse_Monitor_Presentation.pptx"

# ----------------------------------------------------------------- palette
NAVY = RGBColor.from_string("16324F")
INK = RGBColor.from_string("1F2A37")
MUTED = RGBColor.from_string("5B6B7C")
FAINT = RGBColor.from_string("8FA0B3")
LINE = RGBColor.from_string("D8E0E9")
PANEL = RGBColor.from_string("F2F5F9")
PANEL2 = RGBColor.from_string("FDF7EE")
WHITE = RGBColor.from_string("FFFFFF")
RED = RGBColor.from_string("C0392B")
REDBG = RGBColor.from_string("FDF1F0")
AMBER = RGBColor.from_string("B45309")
TEAL = RGBColor.from_string("0F766E")
CYAN = RGBColor.from_string("0E7490")
BLUE = RGBColor.from_string("1D4ED8")
VIOLET = RGBColor.from_string("7C3AED")
INDIGO = RGBColor.from_string("4338CA")
GREEN = RGBColor.from_string("15803D")
LIGHT = RGBColor.from_string("C9D6E4")
MID = RGBColor.from_string("9FB4C9")

FONT = "Segoe UI"
MONO = "Consolas"

SW, SH = 13.333, 7.5          # slide size, inches
ML, MR = 0.55, 0.55           # margins
CW = SW - ML - MR             # content width


def solid(shape, color, line_color=None, line_w=None):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    if line_color is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line_color
        shape.line.width = Pt(line_w or 1)
    shape.shadow.inherit = False


def textbox(slide, x, y, w, h):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return tb, tf


def set_para(p, text, size, color=INK, bold=False, font=FONT, italic=False,
             align=PP_ALIGN.LEFT, space_after=0, line_spacing=1.0):
    p.text = text
    p.alignment = align
    p.space_after = Pt(space_after)
    p.line_spacing = line_spacing
    for r in p.runs:
        r.font.size = Pt(size)
        r.font.color.rgb = color
        r.font.bold = bold
        r.font.italic = italic
        r.font.name = font
    return p


def _bullet(p, char, marL, indent):
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(Inches(marL)))
    pPr.set("indent", str(Inches(indent)))
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    bf = etree.SubElement(pPr, qn("a:buFont"))
    bf.set("typeface", FONT)
    bc = etree.SubElement(pPr, qn("a:buChar"))
    bc.set("char", char)


def bullets(slide, items, x, y, w, h, size=13.5, sub_size=12, space_after=7,
            line_spacing=1.05, numbered=False):
    """items: list of str, or (text, 'sub'). numbered=True renders without bullet chars."""
    tb, tf = textbox(slide, x, y, w, h)
    first = True
    for item in items:
        sub = False
        if isinstance(item, tuple):
            text, kind = item
            sub = kind == "sub"
        else:
            text = item
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        if sub:
            set_para(p, text, sub_size, MUTED, space_after=space_after,
                     line_spacing=line_spacing)
            _bullet(p, "\u2013", 0.52, -0.24)
        elif numbered:
            set_para(p, text, size, INK, space_after=space_after,
                     line_spacing=line_spacing)
        else:
            set_para(p, text, size, INK, space_after=space_after,
                     line_spacing=line_spacing)
            _bullet(p, "\u2022", 0.26, -0.26)
    return tb


def callout(slide, x, y, w, h, text, size=11.5, fill=REDBG, border=RED,
            color=RED, bold=False):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x),
                                Inches(y), Inches(w), Inches(h))
    sh.adjustments[0] = 0.14
    solid(sh, fill, border, 1.2)
    tf = sh.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.12)
    tf.margin_top = tf.margin_bottom = Inches(0.06)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    set_para(tf.paragraphs[0], text, size, color, bold=bold, line_spacing=1.08)
    return sh


def codebox(slide, x, y, w, h, lines, size=11):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x),
                                Inches(y), Inches(w), Inches(h))
    sh.adjustments[0] = 0.10
    solid(sh, RGBColor.from_string("10233B"))
    tf = sh.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.16)
    tf.margin_top = tf.margin_bottom = Inches(0.08)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        set_para(p, ln, size, RGBColor.from_string("D8E6F2"), font=MONO,
                 line_spacing=1.1)
    return sh


def pic(slide, name, x, y, w=None, h=None):
    path = ASSETS / name
    kw = {}
    if w is not None:
        kw["width"] = Inches(w)
    if h is not None:
        kw["height"] = Inches(h)
    return slide.shapes.add_picture(str(path), Inches(x), Inches(y), **kw)


def caption(slide, x, y, w, text):
    tb, tf = textbox(slide, x, y, w, 0.3)
    set_para(tf.paragraphs[0], text, 9, FAINT, italic=True)


def chip_row(slide, x, y, h, chips, gap=0.12, size=10.5):
    cx = x
    for text, color in chips:
        w = 0.14 + 0.085 * len(text)
        sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(cx),
                                    Inches(y), Inches(w), Inches(h))
        sh.adjustments[0] = 0.5
        solid(sh, WHITE, color, 1.1)
        tf = sh.text_frame
        tf.word_wrap = False
        tf.margin_left = tf.margin_right = Inches(0.05)
        tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        set_para(tf.paragraphs[0], text, size, color, align=PP_ALIGN.CENTER)
        cx += w + gap
    return cx


def table(slide, rows_data, x, y, w, col_w, row_h=0.42, header=True,
          font_size=11, header_size=11.5):
    nrows = len(rows_data)
    ncols = len(rows_data[0])
    gf = slide.shapes.add_table(nrows, ncols, Inches(x), Inches(y),
                                Inches(w), Inches(row_h * nrows))
    tbl = gf.table
    for i, cw in enumerate(col_w):
        tbl.columns[i].width = Inches(cw)
    for ri, row in enumerate(rows_data):
        tbl.rows[ri].height = Inches(row_h)
        for ci, val in enumerate(row):
            cell = tbl.cell(ri, ci)
            cell.margin_left = cell.margin_right = Inches(0.09)
            cell.margin_top = cell.margin_bottom = Inches(0.02)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            if header and ri == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = NAVY
                color, bold, size = WHITE, True, header_size
            else:
                cell.fill.solid()
                cell.fill.fore_color.rgb = WHITE if (ri % 2 == 1) else PANEL
                color, bold, size = INK, ci == 0 and False, font_size
            p = cell.text_frame.paragraphs[0]
            set_para(p, str(val), size, color, bold=bold, line_spacing=1.0)
    return tbl


# --------------------------------------------------------------- scaffolding
prs = Presentation()
prs.slide_width = Inches(SW)
prs.slide_height = Inches(SH)
BLANK = prs.slide_layouts[6]
N_SLIDES = 18


def new_slide(title=None):
    slide = prs.slides.add_slide(BLANK)
    if title:
        bar = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(ML),
                                     Inches(0.42), Inches(0.09), Inches(0.5))
        bar.adjustments[0] = 0.5
        solid(bar, RED)
        tb, tf = textbox(slide, ML + 0.26, 0.33, CW - 0.3, 0.7)
        set_para(tf.paragraphs[0], title, 24, NAVY, bold=True)
        ln = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(ML),
                                    Inches(1.08), Inches(CW), Inches(0.016))
        solid(ln, LINE)
    return slide


def footer(slide, n):
    tb, tf = textbox(slide, ML, 7.13, CW - 0.6, 0.3)
    set_para(tf.paragraphs[0], "Pulse Monitor  \u00b7  ESP32 heart-rate & SpO\u2082 pipeline with AI-assisted review",
             8.5, FAINT)
    tb2, tf2 = textbox(slide, SW - MR - 0.9, 7.13, 0.9, 0.3)
    set_para(tf2.paragraphs[0], f"{n} / {N_SLIDES}", 8.5, FAINT, align=PP_ALIGN.RIGHT)


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


# ================================================================== 1 title
s = prs.slides.add_slide(BLANK)
bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(SW), Inches(SH))
solid(bg, NAVY)
acc = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9), Inches(1.42),
                         Inches(1.15), Inches(0.07))
acc.adjustments[0] = 0.5
solid(acc, RED)
tb, tf = textbox(s, 0.9, 0.95, 11, 0.4)
set_para(tf.paragraphs[0], "ACADEMIC PROJECT PRESENTATION", 12.5, LIGHT, bold=True)
tb, tf = textbox(s, 0.9, 1.72, 11.6, 1.1)
set_para(tf.paragraphs[0], "Pulse Monitor", 48, WHITE, bold=True)
tb, tf = textbox(s, 0.9, 2.78, 11.4, 0.9)
set_para(tf.paragraphs[0],
         "ESP32-based heart-rate & SpO\u2082 monitoring with cloud logging "
         "and AI-assisted measurement review", 19, LIGHT, line_spacing=1.15)
tb, tf = textbox(s, 0.9, 3.66, 11.6, 0.4)
set_para(tf.paragraphs[0],
         "Wokwi simulation  \u00b7  ThingSpeak cloud  \u00b7  Python bridge  \u00b7  "
         "AI service  \u00b7  SQLite  \u00b7  web dashboard", 12.5, MID)
div = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.92), Inches(4.32),
                         Inches(4.6), Inches(0.016))
solid(div, RGBColor.from_string("35506E"))
for i, line in enumerate([
        "Presented by:  ______________________",
        "Guide:  ______________________",
        "Department of ______________________   \u00b7   Oct 2026"]):
    tb, tf = textbox(s, 0.9, 4.52 + i * 0.36, 8, 0.34)
    set_para(tf.paragraphs[0], line, 13, LIGHT)
pic(s, "pulse_line.png", 3.4, 5.92, w=9.2)
notes(s, "Greeting + one-line hook: this project is a complete monitoring pipeline \u2014 "
         "from a simulator, through the cloud, to an AI-assisted review dashboard. "
         "Fill in your name, guide and department before presenting.")

# ================================================================== 2 agenda
s = new_slide("Agenda")
bullets(s, [
    "1.  Problem & objectives",
    "2.  System architecture at a glance",
    "3.  Components \u2014 role, objective & how each works",
    "4.  Cloud path: ThingSpeak \u2192 bridge \u2192 local API",
    "5.  What the AI service does \u2014 and what it returned",
    "6.  Dashboard & end-to-end walkthrough of one reading",
    "7.  Verification results",
    "8.  Limitations, safety & future work",
], ML, 1.55, 7.0, 5.2, size=15.5, space_after=11, numbered=True)
box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.15), Inches(1.6),
                         Inches(CW - 7.6), Inches(3.1))
box.adjustments[0] = 0.06
solid(box, PANEL, LINE, 1)
tf = box.text_frame
tf.word_wrap = True
tf.margin_left = tf.margin_right = Inches(0.24)
tf.margin_top = tf.margin_bottom = Inches(0.16)
set_para(tf.paragraphs[0], "At a glance", 14.5, NAVY, bold=True, space_after=8)
for t in ["1 simulation source \u00b7 1 cloud log",
          "3 local services \u00b7 1 database",
          "2-stage AI review layer",
          "every claim verified on the running system"]:
    p = tf.add_paragraph()
    set_para(p, t, 12, MUTED, space_after=6)
footer(s, 2)
notes(s, "Walk through the agenda quickly: architecture first, then each component with "
         "its objective and working, then the AI results, then results and limits.")

# ========================================================= 3 problem/objective
s = new_slide("Why this project: problem & objectives")
tb, tf = textbox(s, ML, 1.4, 5.9, 0.4)
set_para(tf.paragraphs[0], "The problem", 15, RED, bold=True)
bullets(s, [
    "Continuous vitals logging is easy to start but hard to finish \u2014 data ends up stranded in one cloud dashboard.",
    "No local record, no review layer, and no context for what the numbers mean.",
    "Early prototypes often imply medical accuracy they do not have.",
], ML, 1.82, 5.9, 3.0, size=12.5, space_after=8)
tb, tf = textbox(s, 6.95, 1.4, 5.85, 0.4)
set_para(tf.paragraphs[0], "Objectives", 15, TEAL, bold=True)
bullets(s, [
    "Acquire pulse + SpO\u2082 estimates optically (MAX30102 on an ESP32).",
    "Reject invalid measurement windows before anything is uploaded.",
    "Persist every valid reading locally (SQLite) and visualise it live.",
    "Let the user log symptoms so readings have context.",
    "Publish to ThingSpeak for cloud storage and MATLAB analysis.",
    "Add an honest AI layer: transparent rules + unsupervised anomaly scoring.",
], 6.95, 1.82, 5.85, 4.0, size=12.5, space_after=6)
callout(s, ML, 6.15, CW, 0.72,
        "Research prototype \u2014 not a medical device. It does not diagnose anything, "
        "and it is not a substitute for professional care.", size=12.5, bold=False)
footer(s, 3)
notes(s, "Emphasise the gap this addresses and set honest expectations right at the start. "
         "The objective list doubles as the project scope.")

# ================================================================ 4 architecture
s = new_slide("System architecture")
pic(s, "arch_diagram.png", 1.72, 1.28, w=9.9)
tb, tf = textbox(s, ML, 6.35, CW, 0.75)
set_para(tf.paragraphs[0],
         "Two ingestion paths converge on one local API: the cloud path (simulation \u2192 ThingSpeak \u2192 bridge) "
         "and a direct LAN path (physical board). Every stored reading is scored by the AI service, and the "
         "dashboard reads only from the local API \u2014 one source of truth.", 11.5, MUTED, line_spacing=1.1)
footer(s, 4)
notes(s, "Trace the flow left to right. Mention the second, direct path for the physical board, "
         "covered by the dashed note under the Wokwi block.")

# ============================================================ 5 components grid
s = new_slide("Components at a glance \u2014 each has one clear job")
tiles = [
    ("MAX30102 sensor", "Optical PPG: red + IR LEDs and a photodiode yield pulse and SpO\u2082 estimates.", AMBER),
    ("ESP32 DevKit", "Reads the sensor over I\u00b2C, validates windows, uploads over Wi-Fi.", AMBER),
    ("Wokwi simulation", "ESP32 + OLED + pots reproduce the firmware without hardware.", AMBER),
    ("ThingSpeak channel", "Cloud log (Fields 1\u20134) + REST read API for the PC.", TEAL),
    ("Bridge (Python)", "Polls the channel, dedupes entries, forwards to the local API.", CYAN),
    ("Node API + SQLite", "Validation, persistence, AI calls, REST for the dashboard.", BLUE),
    ("AI service (FastAPI)", "Review rules + Isolation Forest anomaly scoring.", VIOLET),
    ("Dashboard (web)", "Metrics, latest AI review, flags table, symptom journal.", INDIGO),
]
tw, th, gx, gy = 2.90, 2.10, 0.21, 0.32
for i, (t, d, c) in enumerate(tiles):
    r, col = divmod(i, 4)
    x = ML + col * (tw + gx)
    y = 1.50 + r * (th + gy)
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y),
                            Inches(tw), Inches(th))
    sh.adjustments[0] = 0.07
    solid(sh, WHITE, c, 1.6)
    tf = sh.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Inches(0.14)
    tf.margin_top = Inches(0.12)
    set_para(tf.paragraphs[0], t, 13, c, bold=True, space_after=5)
    p = tf.add_paragraph()
    set_para(p, d, 10.5, MUTED, line_spacing=1.12)
footer(s, 5)
notes(s, "One sentence per tile; the next slides explain the important ones in depth.")

# ============================================================= 6 core hardware
s = new_slide("Core hardware \u2014 MAX30102 sensor & ESP32")
bullets(s, [
    "MAX30102 \u2014 objective: acquire the pulse waveform optically.",
    ("Red + IR LEDs shine into the skin; a photodiode measures the reflected light.", "sub"),
    ("Pulse rate and SpO\u2082 are estimated from the light ratios (PPG).", "sub"),
    ("It is not an ECG \u2014 beat shape / electrical activity are not captured.", "sub"),
    "ESP32 \u2014 objective: turn raw samples into trustworthy readings.",
    ("Reads the sensor over I\u00b2C in 100-sample windows (SparkFun MAX3010x library).", "sub"),
    ("Uploads only when a window is valid: pulse and SpO\u2082 found, ranges plausible.", "sub"),
    ("Adds a signal-quality heuristic and a beat-interval variability flag \u2014 explicitly not arrhythmia detection.", "sub"),
], ML, 1.45, 6.1, 5.5, size=13, sub_size=11.5, space_after=5)
pic(s, "esp32_photo.jpg", 6.85, 1.5, w=3.15)
caption(s, 6.85, 3.94, 3.4, "ESP32 DevKit (photo: Wikimedia Commons, CC BY-SA 4.0)")
pic(s, "max30102_module.png", 6.85, 4.32, w=2.35)
caption(s, 6.85, 6.12, 4.5, "MAX30102 module \u2014 red + IR LEDs and photodiode (illustration)")
footer(s, 6)
notes(s, "Hold the sensor gently against the fingertip. Stress the difference between PPG and ECG, "
         "and that the firmware's variability flag is a coarse heuristic, not a diagnosis.")

# ================================================================== 7 wiring
s = new_slide("Interface \u2014 wiring the sensor to the board")
pic(s, "wiring.png", ML, 1.55, w=6.45)
bullets(s, [
    "I\u00b2C at 3.3 V: SDA \u2192 GPIO21, SCL \u2192 GPIO22; sensor address 0x57.",
    "The same wiring is mirrored one-to-one inside the Wokwi simulation.",
    "INT pin exists but is not used in this prototype.",
    "Practical quality note: short wires, gentle contact with the fingertip, keep still during measurement.",
], 7.25, 1.6, 5.55, 3.6, size=13, space_after=10)
callout(s, 7.25, 4.35, 5.53, 1.15,
        "Powering the module from 5 V or swapping SDA/SCL are the two most common wiring "
        "mistakes \u2014 the sensor simply returns zeros when it is unhappy.", size=11.5,
        fill=PANEL2, border=AMBER, color=AMBER)
footer(s, 7)
notes(s, "Offer to show the physical wiring. Three points: 3.3 V only, address 0x57, keep still.")

# ================================================================== 8 wokwi
s = new_slide("The Wokwi simulation \u2014 developing without hardware")
bullets(s, [
    "Objective: exercise the complete cloud + AI + dashboard pipeline before real hardware exists.",
    "Contains: ESP32 + SSD1306 OLED + two potentiometers + one push button.",
    "Posts to ThingSpeak every ~20 s, respecting the free-tier limit (\u2265 15 s).",
    "The OLED shows live HR, SpO\u2082, quality and Wi-Fi status during the run.",
], ML, 1.5, 6.4, 4.2, size=13, space_after=10)
table(s, [
    ["Control", "Effect"],
    ["HR potentiometer", "heart rate 45\u2013150 bpm \u2192 Field 1"],
    ["SpO\u2082 potentiometer", "85\u2013100 % \u2192 Field 2"],
    ["Push button", "manual / test flag \u2192 Field 4"],
    ["OLED display", "live status on the virtual device"],
], 7.1, 1.5, 5.68, [1.9, 3.78], row_h=0.52, font_size=11)
callout(s, 7.1, 4.5, 5.68, 1.05,
        "Honest note: the simulation validates the software pipeline \u2014 not sensor accuracy. "
        "Its values are synthetic by design.", size=11.5, fill=PANEL2, border=AMBER, color=AMBER)
footer(s, 8)
notes(s, "Explain how the simulator replaced hardware during development and what it does and does "
         "not prove. Mention the 20 s cadence and why (free-tier rate limit).")

# ============================================================== 9 thingspeak
s = new_slide("Cloud logging \u2014 ThingSpeak channel")
table(s, [
    ["Field", "Meaning", "Values"],
    ["Field 1", "heart rate", "bpm"],
    ["Field 2", "SpO\u2082", "%"],
    ["Field 3", "signal quality", "0\u2013100 %"],
    ["Field 4", "manual / test flag", "0 or 1"],
], ML, 1.55, 6.2, [1.3, 2.9, 2.0], row_h=0.5, font_size=11.5)
bullets(s, [
    "Objective: durable cloud storage, MATLAB visualisations, and a pull point the PC can read.",
    "Working: every accepted reading becomes one timestamped entry; the read API (feeds.json) returns them as JSON.",
    "Only the newest entries matter for live use \u2014 the bridge reads the last 100 and picks up where it left off.",
    "The write key lives only in the firmware; the channel is public for this project.",
], 7.15, 1.55, 5.63, 4.4, size=12.5, space_after=9)
footer(s, 9)
notes(s, "ThingSpeak provides the storage layer and MATLAB plotting. The read API is what the bridge "
         "consumes \u2014 this decouples the cloud from the local network.")

# ================================================================== 10 bridge
s = new_slide("ThingSpeak \u2192 local API bridge (Python)")
bullets(s, [
    "Polls feeds.json every 10 s and sorts entries by ID.",
    "Skips everything already processed \u2014 the state file remembers the last entry ID, so no duplicates and no silent replays.",
    "Validates numeric fields; a broken entry is retried, never skipped.",
    "Forwards each new entry to POST /api/readings and prints the AI result to the console.",
], ML, 1.5, 6.55, 3.4, size=12.5, space_after=8)
codebox(s, ML, 4.62, 6.55, 1.06,
        ["ThingSpeak entry 3 forwarded  |  HR=45 bpm  SpO2=100%",
         "API saved=True  |  AI flags=['heart_rate_outside_configured_review_band']"])
tb, tf = textbox(s, ML, 5.85, 6.55, 1.1)
set_para(tf.paragraphs[0],
         "Config:  THINGSPEAK_CHANNEL_ID (+ optional READ_API_KEY)  \u00b7  first run processes only "
         "the latest entry; full history replay is opt-in via REPLAY_HISTORY_ON_FIRST_RUN.",
         10.5, MUTED, line_spacing=1.15)
callout(s, 7.35, 1.5, 5.43, 1.55,
        "Why this exists: Wokwi runs in the cloud and cannot reach 127.0.0.1 on your PC. "
        "So instead of the device pushing to the laptop, the laptop pulls from the cloud.",
        size=11.5, fill=PANEL2, border=AMBER, color=AMBER)
bullets(s, [
    "Poll the channel (\u2264 100 entries).",
    "Keep only IDs newer than the state file.",
    "Forward to the local API and check saved=True.",
    "Advance the state file only after a successful save.",
], 7.35, 3.35, 5.43, 2.9, size=12, space_after=7)
footer(s, 10)
notes(s, "Explain the localhost problem in one sentence, then the state file as the exactly-once "
         "mechanism. The code box shows real console output from the live run with the user's channel.")

# ================================================================ 11 node api
s = new_slide("Local server \u2014 Node API + SQLite")
table(s, [
    ["Endpoint", "Job"],
    ["POST /api/readings", "validate \u2192 store \u2192 AI score \u2192 save flags"],
    ["GET /api/readings", "recent readings with their AI flags"],
    ["POST/GET /api/symptoms", "symptom journal entries"],
    ["GET /api/health", "service status"],
], ML, 1.55, 6.6, [2.5, 4.1], row_h=0.55, font_size=11)
bullets(s, [
    "Range validation (HR 25\u2013240, SpO\u2082 70\u2013100, quality 0\u2013100) \u2014 bad data never reaches the database.",
    "If AI_SERVICE_URL is configured, every reading is scored and its flags stored alongside.",
    "SQLite in WAL mode with two tables (readings, symptoms); the database is created automatically and older files upgrade in place.",
    "Bound to the local network only \u2014 nothing is exposed to the internet.",
], 7.4, 1.55, 5.38, 4.4, size=12.5, space_after=9)
footer(s, 11)
notes(s, "Point at validation as the first line of data hygiene. Mention SQLite WAL and automatic "
         "schema setup \u2014 zero manual database work.")

# ================================================================ 12 ai service
s = new_slide("What the AI service does")
pic(s, "ai_pipeline.png", 2.05, 1.3, w=9.25)
bullets(s, [
    "A local FastAPI microservice (POST /score) \u2014 no health data leaves the computer.",
    "Stage 1 rules are deterministic and explainable; Stage 2 learns what is normal for this device only, in memory, and needs \u2265 50 samples before it says anything.",
    "The response itself carries the disclaimer: experimental screening aid, not validated for diagnosis or arrhythmia detection.",
], ML, 6.05, CW, 1.0, size=11.5, space_after=4)
footer(s, 12)
notes(s, "Explain the two stages in one breath each: rules anyone can audit; a model that learns "
         "your device's normal. Emphasise the model needs history and resets on restart.")

# ================================================================ 13 ai results
s = new_slide("What the AI returned \u2014 real results")
bullets(s, [
    "Live example \u2014 ThingSpeak entry 3:  HR 45 bpm \u2192  flag: \u201cheart rate outside configured review band\u201d; stored and shown on the dashboard.",
    "Verification session (61 readings):  one HR 152 bpm outlier received 5 flags, including \u201cunusual relative to recent device history\u201d (Isolation Forest, anomaly = 1).",
    "Every flag is a prompt to review context \u2014 never a conclusion.",
], ML, 1.55, 5.15, 2.9, size=12.5, space_after=10)
codebox(s, ML, 4.55, 5.15, 1.8, [
    "\u201creviewFlags\u201d: [",
    "   \u201clow_signal_quality\u201d,",
    "   \u201cheart_rate_outside_configured_review_band\u201d,",
    "   \u201cspo2_below_configured_review_threshold\u201d,",
    "   \u201cpulse_interval_variability_flag_from_firmware\u201d,",
    "   \u201cunusual_relative_to_recent_device_history\u201d",
    "],  \u201cunsupervisedAnomaly\u201d: 1",
], size=8.5)
caption(s, ML, 6.42, 5.15, "The stored AI response for the outlier \u2014 verbatim.")
pic(s, "anomaly_chart.png", 5.95, 1.62, w=6.83)
caption(s, 5.95, 4.78, 6.8, "Verification run: baseline, planned special readings and the flagged outlier.")
footer(s, 13)
notes(s, "Show the real numbers: HR 45 flagged live; the 152 outlier after 61 samples. "
         "Reiterate that flags are advisory only.")

# ================================================================ 14 dashboard
s = new_slide("The dashboard \u2014 where the user meets the data")
pic(s, "dashboard_top.png", ML, 1.42, w=6.5)
bullets(s, [
    "Metric cards: latest heart rate, SpO\u2082, signal quality, sample count.",
    "\u201cLatest AI-assisted review\u201d panel: plain-language flags for the newest reading.",
    "Measurements table: every reading with its AI flags \u2014 plus the symptom journal below.",
    "Auto-refreshes every 10 s; open it from any device on the LAN.",
    "Disclaimers are visible in the interface itself, not buried in documentation.",
], 7.35, 1.5, 5.43, 4.6, size=12.5, space_after=9)
footer(s, 14)
notes(s, "Open the live dashboard if possible. Point at the metric cards, the AI review panel, "
         "and the flags column in the table.")

# =================================================================== 15 e2e
s = new_slide("One reading, end to end")
pic(s, "dataflow.png", ML, 1.35, w=CW)
strip = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(ML), Inches(4.72),
                           Inches(CW), Inches(1.28))
strip.adjustments[0] = 0.09
solid(strip, PANEL, LINE, 1)
tf = strip.text_frame
tf.word_wrap = True
tf.margin_left = tf.margin_right = Inches(0.22)
tf.margin_top = tf.margin_bottom = Inches(0.1)
set_para(tf.paragraphs[0],
         "entry 3  \u00b7  HR 45.0 bpm  \u00b7  SpO\u2082 100 %  \u00b7  quality 95   \u2192   forwarded by the bridge   \u2192   "
         "saved by the API   \u2192   AI flag: heart_rate_outside_configured_review_band   \u2192   "
         "dashboard shows the row with the flag in plain language.", 12, INK, line_spacing=1.15)
bullets(s, [
    "If the API call fails, the bridge retries \u2014 the entry is never marked processed until the save succeeds.",
    "The AI model's history is kept in memory, so a service restart simply starts learning again.",
], ML, 6.25, CW, 0.8, size=11.5, space_after=4)
footer(s, 15)
notes(s, "Follow one real reading end-to-end and land on the dashboard row. This proves the chain "
         "with a real ThingSpeak entry.")

# ================================================================= 16 results
s = new_slide("Verification results \u2014 observed on the running system")
table(s, [
    ["Test", "Result"],
    ["Live chain: ThingSpeak entry \u2192 dashboard", "one real reading stored, scored and displayed \u2014 end to end in under one poll cycle  \u2713"],
    ["Bridge duplicate safety", "restarts and re-polls produced no duplicate rows  \u2713"],
    ["Broken-entry handling", "corrupt entries were retried; the state file never advanced past unsaved data  \u2713"],
    ["First-run policy", "only the latest existing entry processed by default; replay is opt-in  \u2713"],
    ["AI anomaly stage", "outlier flagged after 61 in-service samples (anomaly = 1)  \u2713"],
    ["Dashboard QA", "0 console errors, 0 failed network requests during the full session  \u2713"],
    ["Integration defects found & fixed", "AI service startup, console encoding, SpO\u2082/quality unit display  \u2713"],
], ML, 1.45, CW, [4.6, 7.63], row_h=0.55, font_size=11)
tb, tf = textbox(s, ML, 6.42, CW, 0.5)
set_para(tf.paragraphs[0],
         "Every row above was observed on the actual pipeline against the project's real ThingSpeak "
         "channel \u2014 not estimated.", 11, MUTED)
footer(s, 16)
notes(s, "Summarise the verification matrix. Every claim here was observed on the running system; "
         "the last row shows engineering hygiene \u2014 defects were found and fixed during testing.")

# ========================================================== 17 limits / future
s = new_slide("Limitations, safety & what comes next")
tb, tf = textbox(s, ML, 1.4, 5.9, 0.4)
set_para(tf.paragraphs[0], "Limitations \u2014 stated honestly", 14.5, RED, bold=True)
bullets(s, [
    "Sensor data in the simulation is synthetic by design.",
    "An optical PPG sensor is not an ECG and cannot see beat morphology.",
    "The anomaly model is unsupervised and in-memory (it resets on restart).",
    "Thresholds are prototype choices, not clinical values; symptoms are self-reported.",
    "Local use only: no accounts, no encryption at rest.",
], ML, 1.82, 5.9, 4.4, size=12, space_after=7)
tb, tf = textbox(s, 7.0, 1.4, 5.78, 0.4)
set_para(tf.paragraphs[0], "Safety rules & next steps", 14.5, GREEN, bold=True)
bullets(s, [
    "Frame it as a logging and review aid \u2014 never as a diagnosis; keep disclaimers visible.",
    "Keep API keys private; rotate any key that is ever shared in a screenshot.",
    "Next: collect real labelled data with an ECG reference.",
    "Then: validated metrics (sensitivity / specificity) and a supervised model trained on that data.",
    "Later: persistent training, alerts, multi-user access, documented protocol.",
], 7.0, 1.82, 5.78, 4.4, size=12, space_after=7)
footer(s, 17)
notes(s, "Be candid about limitations; frame next steps as the path to research-grade evidence.")

# ============================================================ 18 conclusion
s = new_slide("Conclusion")
bullets(s, [
    "A complete pipeline was built and verified end to end: simulation \u2192 cloud \u2192 bridge \u2192 API \u2192 AI review \u2192 SQLite \u2192 dashboard.",
    "Each component has one clear job, and every failure mode found during testing was fixed and re-verified.",
    "The AI layer adds transparent review flags plus an experimental anomaly signal \u2014 framed as an aid, not a diagnosis.",
    "The system is ready for its next phase: real sensor data and formal validation.",
], ML, 1.55, CW - 0.4, 3.4, size=14.5, space_after=13)
callout(s, ML, 5.35, CW, 1.15, "Thank you \u2014 questions?", size=22, fill=PANEL,
        border=NAVY, color=NAVY, bold=True)
footer(s, 18)
notes(s, "Close with the one-liner: the pipeline works end to end and is honest about what it is. "
         "Invite questions.")

# ------------------------------------------------------------------- save
prs.core_properties.title = "Pulse Monitor \u2014 ESP32 heart-rate & SpO2 pipeline with AI-assisted review"
prs.core_properties.author = "Pulse Monitor project"
prs.core_properties.subject = "Academic project presentation"
prs.save(OUT)
print(f"wrote {OUT}  ({len(prs.slides._sldIdLst)} slides)")
